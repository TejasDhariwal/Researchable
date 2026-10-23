from pptx import Presentation
from pptx.util import Pt

def create_presentation(title,pages):
    prs = Presentation()

    title_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_layout)
    Title = title_slide.shapes.title
    Title.text = title


    for page in pages:
        slide_layout = prs.slide_layouts[1] # using title and content layout
        slide = prs.slides.add_slide(slide_layout)

        # extract title and content sections from the slide
        shapes = slide.shapes

        title_shape = shapes.title
        title_shape.text_frame.paragraphs[0].font.size = Pt(32)
        title_shape.text = page["title"]
        
        # Add each list element as a bullet point
        text_frame = shapes.placeholders[1].text_frame
        text_frame.clear()

        for i, point in enumerate(page["key_points"]):
            if i==0:
                paragraph = text_frame.paragraphs[0]
            else:
                paragraph = text_frame.add_paragraph()

            paragraph.text = point
            paragraph.font.size = Pt(22)
            paragraph.level = 0  # set the indentation level for bullet points

        # Add a heading for the sources
        # paragraph = text_frame.add_paragraph()
        # paragraph.text = "Sources:"
        # paragraph.level = 0
        # paragraph.font.bold = True

        # Add each source URL as a bullet point
        # for url in page["source_urls"]:
        #     paragraph = text_frame.add_paragraph()
        #     paragraph.text = url
        #     paragraph.level = 1
        #     paragraph.font.size = Pt(15)
        
        # add speaker notes
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = page['speaker_notes']
    
    prs.save('presentation.pptx')
